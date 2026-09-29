"""
Generate data dummy PMB (Penerimaan Mahasiswa Baru):
- Sheet 'calon_mahasiswa' : 20 data simulasi calon mahasiswa
- Sheet 'program_studi'   : 5 data simulasi program studi

Output: data_pmb.xlsx
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# DATA PROGRAM STUDI (5 data)
# ---------------------------------------------------------------------------
program_studi = [
    # kode, nama, jenjang, fakultas, akreditasi, kuota
    ["TIF", "Teknik Informatika",     "S1", "Fakultas Teknik",            "A",           60],
    ["SIF", "Sistem Informasi",       "S1", "Fakultas Teknik",            "B",           50],
    ["MNJ", "Manajemen",              "S1", "Fakultas Ekonomi & Bisnis",  "A",           80],
    ["AKT", "Akuntansi",              "S1", "Fakultas Ekonomi & Bisnis",  "Unggul",      70],
    ["DKV", "Desain Komunikasi Visual","S1","Fakultas Seni & Desain",     "B",           40],
]

# ---------------------------------------------------------------------------
# DATA CALON MAHASISWA (20 data)
# ---------------------------------------------------------------------------
calon_mahasiswa = [
    # no_pendaftaran, nama, jenis_kelamin, ttl, asal_sekolah, jurusan_sekolah, email, no_hp, alamat, pilihan_prodi
    ["PMB2026001", "Ahmad Fauzan Rahman",   "Laki-laki",  "Bandung, 2007-03-12",   "SMAN 3 Bandung",       "IPA", "ahmad.fauzan@gmail.com",    "081234567001", "Jl. Merdeka No. 12, Bandung",       "TIF"],
    ["PMB2026002", "Siti Nurhaliza",        "Perempuan",  "Jakarta, 2007-07-25",   "SMAN 8 Jakarta",       "IPA", "siti.nurhaliza@gmail.com",  "081234567002", "Jl. Sudirman No. 45, Jakarta",     "SIF"],
    ["PMB2026003", "Budi Santoso",          "Laki-laki",  "Surabaya, 2006-11-30",  "SMAN 5 Surabaya",      "IPS", "budi.santoso@gmail.com",    "081234567003", "Jl. Pemuda No. 8, Surabaya",       "MNJ"],
    ["PMB2026004", "Dewi Lestari",          "Perempuan",  "Yogyakarta, 2007-01-18","SMAN 1 Yogyakarta",    "IPA", "dewi.lestari@gmail.com",    "081234567004", "Jl. Malioboro No. 20, Yogyakarta", "AKT"],
    ["PMB2026005", "Rizky Pratama",         "Laki-laki",  "Medan, 2007-05-09",     "SMA Methodist Medan",  "IPA", "rizky.pratama@gmail.com",   "081234567005", "Jl. Gatot Subroto No. 3, Medan",   "TIF"],
    ["PMB2026006", "Putri Ayu Wulandari",   "Perempuan",  "Semarang, 2007-09-14",  "SMAN 2 Semarang",      "IPS", "putri.ayu@gmail.com",       "081234567006", "Jl. Pandanaran No. 17, Semarang",  "DKV"],
    ["PMB2026007", "Muhammad Iqbal",        "Laki-laki",  "Makassar, 2006-12-22",  "SMAN 1 Makassar",      "IPA", "m.iqbal@gmail.com",         "081234567007", "Jl. Perintis No. 9, Makassar",     "SIF"],
    ["PMB2026008", "Nabila Zahra",          "Perempuan",  "Padang, 2007-02-28",    "SMAN 10 Padang",       "IPA", "nabila.zahra@gmail.com",    "081234567008", "Jl. Diponegoro No. 5, Padang",     "AKT"],
    ["PMB2026009", "Fajar Nugroho",         "Laki-laki",  "Malang, 2007-08-03",    "SMAN 4 Malang",        "IPS", "fajar.nugroho@gmail.com",   "081234567009", "Jl. Ijen No. 22, Malang",          "MNJ"],
    ["PMB2026010", "Anisa Rahmawati",       "Perempuan",  "Bogor, 2007-06-11",     "SMAN 1 Bogor",         "IPA", "anisa.rahma@gmail.com",     "081234567010", "Jl. Pajajaran No. 30, Bogor",      "TIF"],
    ["PMB2026011", "Dimas Aditya",          "Laki-laki",  "Denpasar, 2006-10-19",  "SMAN 3 Denpasar",      "IPA", "dimas.aditya@gmail.com",    "081234567011", "Jl. Gajah Mada No. 7, Denpasar",   "SIF"],
    ["PMB2026012", "Rani Kusuma",           "Perempuan",  "Palembang, 2007-04-07", "SMAN 6 Palembang",     "IPS", "rani.kusuma@gmail.com",     "081234567012", "Jl. Sudirman No. 14, Palembang",   "DKV"],
    ["PMB2026013", "Yoga Prasetyo",         "Laki-laki",  "Solo, 2007-03-27",      "SMAN 1 Surakarta",     "IPA", "yoga.prasetyo@gmail.com",   "081234567013", "Jl. Slamet Riyadi No. 2, Solo",    "TIF"],
    ["PMB2026014", "Intan Permata",         "Perempuan",  "Balikpapan, 2007-07-16","SMAN 2 Balikpapan",    "IPA", "intan.permata@gmail.com",   "081234567014", "Jl. Ahmad Yani No. 11, Balikpapan","AKT"],
    ["PMB2026015", "Hendra Wijaya",         "Laki-laki",  "Bandung, 2006-09-05",   "SMA BPI Bandung",      "IPS", "hendra.wijaya@gmail.com",   "081234567015", "Jl. Asia Afrika No. 40, Bandung",  "MNJ"],
    ["PMB2026016", "Salsabila Putri",       "Perempuan",  "Jakarta, 2007-11-21",   "SMAN 70 Jakarta",      "IPA", "salsabila.p@gmail.com",     "081234567016", "Jl. Kebayoran No. 6, Jakarta",     "SIF"],
    ["PMB2026017", "Andre Kurniawan",       "Laki-laki",  "Pekanbaru, 2007-01-08", "SMAN 8 Pekanbaru",     "IPA", "andre.kurniawan@gmail.com", "081234567017", "Jl. Sudirman No. 25, Pekanbaru",   "TIF"],
    ["PMB2026018", "Melati Anggraini",      "Perempuan",  "Manado, 2007-05-30",    "SMAN 1 Manado",        "IPS", "melati.a@gmail.com",        "081234567018", "Jl. Sam Ratulangi No. 3, Manado",  "DKV"],
    ["PMB2026019", "Reza Firmansyah",       "Laki-laki",  "Bekasi, 2006-12-14",    "SMAN 1 Bekasi",        "IPA", "reza.firmansyah@gmail.com", "081234567019", "Jl. Ahmad Yani No. 18, Bekasi",    "AKT"],
    ["PMB2026020", "Cantika Dewi",          "Perempuan",  "Tangerang, 2007-08-27", "SMAN 2 Tangerang",     "IPA", "cantika.dewi@gmail.com",    "081234567020", "Jl. MH Thamrin No. 9, Tangerang",  "MNJ"],
]

# ---------------------------------------------------------------------------
# STYLING HELPERS
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_sheet(ws, headers, rows, widths):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    for row in rows:
        ws.append(row)
    for r in range(2, len(rows) + 2):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = LEFT
            cell.border = BORDER
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# BUILD WORKBOOK
# ---------------------------------------------------------------------------
wb = Workbook()

ws1 = wb.active
ws1.title = "calon_mahasiswa"
style_sheet(
    ws1,
    ["No. Pendaftaran", "Nama Lengkap", "Jenis Kelamin", "Tempat, Tgl Lahir",
     "Asal Sekolah", "Jurusan", "Email", "No. HP", "Alamat", "Pilihan Prodi"],
    calon_mahasiswa,
    [16, 24, 14, 22, 22, 10, 26, 16, 34, 14],
)

ws2 = wb.create_sheet("program_studi")
style_sheet(
    ws2,
    ["Kode Prodi", "Nama Program Studi", "Jenjang", "Fakultas", "Akreditasi", "Kuota"],
    program_studi,
    [12, 28, 10, 30, 12, 10],
)

wb.save("data_pmb.xlsx")
print("OK -> data_pmb.xlsx")
print(f"calon_mahasiswa: {len(calon_mahasiswa)} baris")
print(f"program_studi  : {len(program_studi)} baris")
